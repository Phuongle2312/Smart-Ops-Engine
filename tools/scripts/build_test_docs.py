#!/usr/bin/env python3
"""Sinh workbook Excel tổng hợp use case + test case từ các file Markdown trong docs/.

Chạy từ gốc repo:  python tools/scripts/build_test_docs.py
Đầu ra:            docs/SOE_UseCase_TestCase.xlsx
Nguồn dữ liệu là Markdown (docs/03_usecases, 04_test_cases, legacy-v1/test_cases, 06_workplan);
sửa Markdown rồi chạy lại — không sửa tay file Excel (trừ cột theo dõi kết quả).
"""
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / "docs"
OUT = DOCS / "SOE_UseCase_TestCase.xlsx"

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
SUB_FILL = PatternFill("solid", fgColor="D9E2F3")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

MODULES = {
    "IDN": "Identity", "INV": "Inventory", "MON": "Monitoring", "MET": "Metrics",
    "INC": "Incident", "NTF": "Notification", "RTM": "Realtime", "AUD": "Audit",
    "GW": "Gateway & Config", "SEC": "Security", "PERF": "Performance", "MSG": "Messaging",
}
LEGACY_MODULES = {
    "AUTH": "01 Xác thực", "NODE": "02 Quản lý node", "SCHED": "03 Health check scheduler",
    "INCIDENT": "04 Quản lý sự cố", "INC": "04 Quản lý sự cố", "METRIC": "08 Lịch sử metrics", "ALERT": "05 Cảnh báo", "WS": "06 WebSocket realtime",
    "AUDIT": "07 Audit log", "METRICS": "08 Lịch sử metrics", "CONFIG": "09 Cấu hình hệ thống",
}
PRIORITY_NAME = {"P1": "P1 - Cao", "P2": "P2 - Trung bình", "P3": "P3 - Thấp"}


def clean(text):
    text = re.sub(r"<br\s*/?>", "\n", text.strip())
    text = text.replace("**", "").replace("\\|", "|")
    return text.strip()


def split_row(line):
    cells = re.split(r"(?<!\\)\|", line.strip())
    return [clean(c) for c in cells[1:-1]]


def is_sep(line):
    return bool(re.match(r"^\|[\s:\-|]+\|$", line.strip()))


def parse_tables(path):
    """Trả về list (heading_section, headers, rows) cho mọi bảng trong file."""
    tables, section, lines = [], "", path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("## "):
            section = re.sub(r"^##\s*\d*\.?\s*", "", line).strip()
        if line.startswith("|") and i + 1 < len(lines) and is_sep(lines[i + 1]):
            headers = split_row(line)
            rows, i = [], i + 2
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            tables.append((section, headers, rows))
            continue
        i += 1
    return tables


COLMAP = {
    "ID": "id", "Kịch bản": "scenario", "Test Scenario": "scenario", "Hạng mục": "scenario",
    "Rủi ro": "scenario", "Điều kiện tiên quyết": "pre", "Prerequisites": "pre",
    "Các bước": "steps", "Steps": "steps", "Cách thực hiện": "steps", "Công cụ": "steps",
    "Kịch bản kiểm thử": "steps", "Dữ liệu": "data", "Test Data": "data",
    "Cấu hình tải": "data", "Kết quả mong đợi": "expected", "Expected Result": "expected",
    "Chỉ tiêu": "expected", "Tiêu chí đạt": "expected", "Loại": "type", "Type": "type",
    "Ưu tiên": "prio", "Truy vết": "trace", "Status": "status", "Tần suất": "note",
}
TYPE_BY_TAG = {
    "API": "API", "INT": "Integration", "E2E": "E2E (UI)", "SEC": "Security", "PERF": "Performance",
    "UNIT": "Unit", "MSG": "Messaging", "FE": "Frontend", "BE": "Backend",
}


def layer_of(tc_id, section):
    for tag, name in TYPE_BY_TAG.items():
        if f"-{tag}-" in tc_id:
            return name
    s = section.lower()
    if "frontend" in s or "giao diện" in s:
        return "Frontend"
    if "backend" in s:
        return "Backend"
    return section or ""


def collect_tests(files, legacy):
    rows = []
    for path in files:
        for section, headers, body in parse_tables(path):
            if "ID" not in headers or not body or not re.match(r"^(TC|SEC|PERF|MSG|CT)-", body[0][0]):
                continue
            keys = [COLMAP.get(h) for h in headers]
            for cells in body:
                rec = {k: "" for k in set(COLMAP.values())}
                for k, v in zip(keys, cells):
                    if k:
                        rec[k] = (rec[k] + "\n" + v).strip() if rec[k] else v
                parts = rec["id"].split("-")
                tag = parts[1] if len(parts) > 1 else ""
                rec["module"] = (LEGACY_MODULES if legacy else MODULES).get(tag, path.stem)
                if not legacy and rec["id"].startswith(("SEC", "PERF", "CT", "MSG")):
                    rec["module"] = MODULES.get(parts[0], "Cross-cutting")
                if not legacy and path.stem in ("security_tests", "performance_scaling_tests",
                                                "contract_messaging_tests"):
                    rec["module"] = {"security_tests": "Security",
                                     "performance_scaling_tests": "Performance & Scaling",
                                     "contract_messaging_tests": "Messaging & Contract"}[path.stem]
                rec["layer"] = layer_of(rec["id"], section)
                rec["group"] = section
                rec["src"] = path.name
                rows.append(rec)
    return rows


def collect_usecases():
    ucs, steps = [], []
    for path in sorted((DOCS / "03_usecases").glob("UC-*.md")):
        text = path.read_text(encoding="utf-8")
        for block in re.split(r"(?m)^## (?=UC-)", text)[1:]:
            head, _, rest = block.partition("\n")
            m = re.match(r"(UC-[A-Z]+-\d+)\s*[—-]\s*(.+)", head.strip())
            uc = {"id": m.group(1), "name": m.group(2), "module": m.group(1).split("-")[1],
                  "src": path.name, "actor": "", "actor2": "", "prio": "", "freq": "", "pre": "",
                  "trigger": "", "post": "", "br": "", "tc": "", "main": [], "alt": [], "exc": []}
            for mm in re.finditer(r"^\| \*\*(.+?)\*\* \| (.+?) \|$", rest, re.M):
                key, val = mm.group(1), clean(mm.group(2))
                uc_field(uc, key, val)
            section = None
            for line in rest.splitlines():
                s = line.strip()
                if re.match(r"^\*\*Luồng chính", s):
                    section = "main"
                elif re.match(r"^\*\*Luồng thay thế", s):
                    section = "alt"
                elif re.match(r"^\*\*Luồng ngoại lệ", s):
                    section = "exc"
                elif s.startswith("**"):
                    section = "post" if s.startswith("**Hậu điều kiện") else None
                    # dòng meta dạng "**BR:** a · **Test case:** b"
                    for seg in re.split(r"\s+·\s+(?=\*\*)", s):
                        mm = re.match(r"\*\*(.+?)\*\*:?\s*:?\s*(.*)$", seg)
                        if not mm:
                            continue
                        key, val = mm.group(1).strip(": ").lower(), clean(mm.group(2))
                        if key.startswith("hậu điều kiện"):
                            uc["post"] = val
                        elif key.startswith("br") or key.startswith("nfr"):
                            uc["br"] = (uc["br"] + "; " if uc["br"] else "") + clean(mm.group(1)).rstrip(": ") + ": " + val
                        elif key.startswith("test case"):
                            uc["tc"] = val
                elif section == "main" and re.match(r"^\d+\.", s):
                    uc["main"].append(re.sub(r"^\d+\.\s*", "", s))
                elif section in ("alt", "exc") and s.startswith("- "):
                    uc[section].append(s[2:])
                elif section == "post" and s.startswith("- "):
                    uc["post"] = (uc["post"] + chr(10) if uc["post"] else "") + clean(s[2:])
            ucs.append(uc)
            for n, t in enumerate(uc["main"], 1):
                steps.append((uc["id"], uc["name"], "Luồng chính", f"B{n}", clean(t)))
            for kind, lab in (("alt", "Luồng thay thế"), ("exc", "Luồng ngoại lệ")):
                for t in uc[kind]:
                    mm = re.match(r"\*\*(.+?)\*\*\s*[:—-]?\s*(.*)", t)
                    code, body = (clean(mm.group(1)), clean(mm.group(2))) if mm else ("", clean(t))
                    title = code.split("—", 1)[1].strip().rstrip(":") if "—" in code else code.rstrip(":")
                    steps.append((uc["id"], uc["name"], lab, code.split("—")[0].strip(),
                                  f"{title}: {body}" if mm and title and body else (body or title)))
    return ucs, steps


def uc_field(uc, key, val):
    """Điền trường từ bảng 'Mục | Nội dung'; một hàng có thể chứa nhiều trường nối bằng ' · '."""
    segs = [(key, val)]
    while True:
        k0, v0 = segs[-1]
        m = re.match(r"^(.*?)\s*·\s*(.+)$", v0)
        if not m:
            break
        segs[-1] = (k0, m.group(1))
        nk, _, nv = m.group(2).partition(":")
        segs.append((nk.strip(), nv.strip()))
    fields = {"actor chính": "actor", "actor phụ": "actor2", "ưu tiên": "prio", "tần suất": "freq",
              "tiền điều kiện": "pre", "kích hoạt": "trigger", "hậu điều kiện": "post"}
    for k, v in segs:
        k = k.strip("* ").lower()
        for prefix, field in fields.items():
            if k.startswith(prefix):
                uc[field] = clean(v)


def collect_workplan():
    plan, debt = [], []
    for section, headers, rows in parse_tables(DOCS / "06_workplan.md"):
        if headers and headers[0].startswith("Hạng mục"):
            plan = [headers] + rows
        elif headers and headers[0] == "#":
            debt = [headers] + rows
    return plan, debt


# ---------------------------------------------------------------- Excel
def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BORDER


def write_table(ws, headers, rows, widths, start=1, freeze_col="B"):
    for c, h in enumerate(headers, 1):
        ws.cell(row=start, column=c, value=h)
    style_header(ws, start, len(headers))
    for r, row in enumerate(rows, start + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(name=FONT, size=9)
            cell.alignment = WRAP
            cell.border = BORDER
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"{freeze_col}{start + 1}"
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(headers))}{start + len(rows)}"
    ws.row_dimensions[start].height = 32
    return start + len(rows)


def tc_sheet(wb, title, tests, legacy):
    ws = wb.create_sheet(title)
    headers = ["STT", "Mã TC", "Phân hệ", "Nhóm / Tầng", "Kịch bản kiểm thử", "Điều kiện tiên quyết",
               "Các bước thực hiện", "Dữ liệu kiểm thử", "Kết quả mong đợi", "Loại", "Ưu tiên",
               "Truy vết (FR/BR/UC)", "Trạng thái", "Kết quả thực tế", "Người test", "Ngày test"]
    rows = []
    for n, t in enumerate(tests, 1):
        status = t["status"].strip("` ") if legacy and t["status"] else ""
        status = {"Pass": "Pass", "Fail": "Fail"}.get(status, "Chưa chạy")
        rows.append([n, t["id"], t["module"], f'{t["layer"]} — {t["group"]}' if t["group"] else t["layer"],
                     t["scenario"], t["pre"], t["steps"], t["data"], t["expected"], t["type"],
                     PRIORITY_NAME.get(t["prio"], t["prio"]), t["trace"], status, "", "", ""])
    last = write_table(ws, headers, rows,
                       [5, 17, 16, 22, 34, 28, 40, 30, 46, 13, 12, 20, 12, 26, 12, 11], freeze_col="C")
    dv = DataValidation(type="list", formula1='"Chưa chạy,Pass,Fail,Blocked,Skip"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"M2:M{last}")
    for r in range(2, last + 1):
        for c in (13, 14, 15, 16):
            ws.cell(row=r, column=c).fill = INPUT_FILL
    green = PatternFill("solid", bgColor="C6EFCE")
    red = PatternFill("solid", bgColor="FFC7CE")
    amber = PatternFill("solid", bgColor="FFEB9C")
    ws.conditional_formatting.add(f"M2:M{last}", CellIsRule(operator="equal", formula=['"Pass"'], fill=green))
    ws.conditional_formatting.add(f"M2:M{last}", CellIsRule(operator="equal", formula=['"Fail"'], fill=red))
    ws.conditional_formatting.add(f"M2:M{last}", CellIsRule(operator="equal", formula=['"Blocked"'], fill=amber))
    return ws, last


def build():
    wb = Workbook()
    wb.remove(wb.active)

    v3_files = sorted((DOCS / "04_test_cases").glob("*.md"))
    v3 = collect_tests([f for f in v3_files if f.name != "README.md"], legacy=False)
    legacy = collect_tests(sorted((DOCS / "legacy-v1" / "test_cases").glob("0*.md")), legacy=True)
    ucs, steps = collect_usecases()
    plan, debt = collect_workplan()

    # ---- Tổng quan (điền sau khi biết số dòng)
    ov = wb.create_sheet("Tổng quan")

    # ---- Use case
    ws = wb.create_sheet("Use case")
    rows = []
    for n, u in enumerate(ucs, 1):
        rows.append([n, u["id"], u["name"], MODULES.get(u["module"], u["module"]), u["actor"], u["actor2"],
                     u["prio"], u["freq"], u["trigger"], u["pre"],
                     "\n".join(f"{i}. {clean(s)}" for i, s in enumerate(u["main"], 1)),
                     "\n".join(f"- {clean(s)}" for s in u["alt"]),
                     "\n".join(f"- {clean(s)}" for s in u["exc"]),
                     u["post"], u["br"], u["tc"], u["src"]])
    write_table(ws, ["STT", "Mã UC", "Tên use case", "Phân hệ", "Actor chính", "Actor phụ", "Ưu tiên",
                     "Tần suất", "Kích hoạt", "Tiền điều kiện", "Luồng chính", "Luồng thay thế",
                     "Luồng ngoại lệ", "Hậu điều kiện", "Business rule / NFR", "Test case liên quan",
                     "Nguồn"], rows,
                [5, 13, 28, 14, 18, 18, 10, 14, 26, 30, 60, 46, 56, 34, 22, 28, 20], freeze_col="D")

    # ---- Các bước use case
    ws = wb.create_sheet("Use case - Từng bước")
    write_table(ws, ["Mã UC", "Tên use case", "Loại luồng", "Mã bước", "Nội dung"], [list(s) for s in steps],
                [13, 30, 16, 10, 110], freeze_col="A")

    # ---- Test case v3
    ws_v3, last_v3 = tc_sheet(wb, "Test case v3", v3, legacy=False)
    ws_lg, last_lg = tc_sheet(wb, "Test case v1 (legacy)", legacy, legacy=True)

    # ---- Ma trận UC ↔ module
    ws = wb.create_sheet("Truy vết")
    hdr = ["Phân hệ", "Số use case", "Số test case v3", "P1", "P2", "P3", "Đã Pass", "Đã Fail", "Chưa chạy"]
    ws.append(hdr)
    style_header(ws, 1, len(hdr))
    names = ["Identity", "Inventory", "Monitoring", "Metrics", "Incident", "Notification", "Realtime",
             "Audit", "Gateway & Config", "Security", "Performance & Scaling", "Messaging & Contract"]
    uc_col = "Use case"
    for r, name in enumerate(names, 2):
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=f"=COUNTIF('{uc_col}'!$D$2:$D$500,A{r})")
        rng = f"'Test case v3'!$C$2:$C${last_v3}"
        ws.cell(row=r, column=3, value=f"=COUNTIF({rng},A{r})")
        for c, p in zip((4, 5, 6), ("P1*", "P2*", "P3*")):
            ws.cell(row=r, column=c, value=f"=COUNTIFS({rng},$A{r},'Test case v3'!$K$2:$K${last_v3},\"{p}\")")
        for c, s in zip((7, 8, 9), ("Pass", "Fail", "Chưa chạy")):
            ws.cell(row=r, column=c, value=f"=COUNTIFS({rng},$A{r},'Test case v3'!$M$2:$M${last_v3},\"{s}\")")
    tot = len(names) + 2
    ws.cell(row=tot, column=1, value="Tổng")
    for c in range(2, 10):
        L = get_column_letter(c)
        ws.cell(row=tot, column=c, value=f"=SUM({L}2:{L}{tot - 1})")
    for r in range(2, tot + 1):
        for c in range(1, 10):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT, size=10, bold=(r == tot))
            cell.border = BORDER
            if r == tot:
                cell.fill = SUB_FILL
    ws.column_dimensions["A"].width = 26
    for c in range(2, 10):
        ws.column_dimensions[get_column_letter(c)].width = 14
    ws.cell(row=tot + 2, column=1,
            value="Ghi chú: 'Security', 'Performance & Scaling', 'Messaging & Contract' là bộ test xuyên service, "
                  "không gắn use case riêng. Xem cột 'Truy vết' ở sheet Test case v3 để biết FR/BR tương ứng.").font = Font(
        name=FONT, italic=True, size=9)

    # ---- Kế hoạch công việc
    if plan:
        ws = wb.create_sheet("Công việc (workplan)")
        end = write_table(ws, plan[0], plan[1:], [34, 44, 34, 34, 8, 16], freeze_col="B")
        if debt:
            start = end + 3
            ws.cell(row=start - 1, column=1, value="Việc đang nợ của M1").font = Font(name=FONT, bold=True, size=11)
            for c, h in enumerate(debt[0], 1):
                ws.cell(row=start, column=c, value=h)
            style_header(ws, start, len(debt[0]))
            for r, row in enumerate(debt[1:], start + 1):
                for c, v in enumerate(row, 1):
                    cell = ws.cell(row=r, column=c, value=v)
                    cell.font = Font(name=FONT, size=9)
                    cell.alignment = WRAP
                    cell.border = BORDER

    # ---- Tổng quan
    ov["A1"] = "Smart Ops Engine — Use case & Test case"
    ov["A1"].font = Font(name=FONT, bold=True, size=16, color="1F3864")
    ov["A2"] = "Sinh tự động từ docs/*.md bằng tools/scripts/build_test_docs.py — sửa Markdown rồi chạy lại."
    ov["A2"].font = Font(name=FONT, italic=True, size=9)
    stats = [
        ("Use case (v3)", f"=COUNTA('Use case'!B2:B500)"),
        ("Bước / luồng use case", f"=COUNTA('Use case - Từng bước'!A2:A2000)"),
        ("Test case v3", f"=COUNTA('Test case v3'!B2:B{last_v3})"),
        ("Test case v1 (legacy)", f"=COUNTA('Test case v1 (legacy)'!B2:B{last_lg})"),
        ("Test case v3 — Pass", f"=COUNTIF('Test case v3'!M2:M{last_v3},\"Pass\")"),
        ("Test case v3 — Fail", f"=COUNTIF('Test case v3'!M2:M{last_v3},\"Fail\")"),
        ("Test case v3 — Chưa chạy", f"=COUNTIF('Test case v3'!M2:M{last_v3},\"Chưa chạy\")"),
        ("Test case v3 — Ưu tiên P1", f"=COUNTIF('Test case v3'!K2:K{last_v3},\"P1*\")"),
        ("Test case v1 — Pass", f"=COUNTIF('Test case v1 (legacy)'!M2:M{last_lg},\"Pass\")"),
        ("Test case v1 — Chưa chạy", f"=COUNTIF('Test case v1 (legacy)'!M2:M{last_lg},\"Chưa chạy\")"),
    ]
    ov["A4"], ov["B4"] = "Chỉ số", "Giá trị"
    style_header(ov, 4, 2)
    for r, (k, f) in enumerate(stats, 5):
        ov.cell(row=r, column=1, value=k).font = Font(name=FONT, size=10)
        c = ov.cell(row=r, column=2, value=f)
        c.font = Font(name=FONT, size=10, bold=True)
        for cc in (ov.cell(row=r, column=1), c):
            cc.border = BORDER
    r = 5 + len(stats) + 1
    ov.cell(row=r, column=1, value="Các sheet").font = Font(name=FONT, bold=True, size=11)
    guide = [
        ("Use case", "Mỗi dòng một use case: actor, kích hoạt, tiền/hậu điều kiện, luồng chính/thay thế/ngoại lệ, BR, test case liên quan."),
        ("Use case - Từng bước", "Mỗi dòng một bước/luồng của use case — dùng để soi chi tiết hoặc lọc theo UC."),
        ("Test case v3", "Test case bản .NET 8 microservices (API, Integration, E2E UI, Security, Performance, Messaging). Ô vàng để ghi kết quả."),
        ("Test case v1 (legacy)", "Test case bản Spring Boot v1 đang chạy (Backend + Frontend theo 9 module)."),
        ("Truy vết", "Thống kê use case / test case / ưu tiên / kết quả theo phân hệ (công thức, tự cập nhật)."),
        ("Công việc (workplan)", "Phân chia hạng mục ↔ thư mục ↔ SRS ↔ test case ↔ mốc M1–M5 và việc nợ của M1."),
    ]
    for i, (k, v) in enumerate(guide, r + 1):
        ov.cell(row=i, column=1, value=k).font = Font(name=FONT, bold=True, size=10)
        ov.cell(row=i, column=2, value=v).font = Font(name=FONT, size=10)
        ov.cell(row=i, column=2).alignment = WRAP
    i = r + len(guide) + 2
    ov.cell(row=i, column=1, value="Cách ghi kết quả").font = Font(name=FONT, bold=True, size=11)
    for j, txt in enumerate([
        "Chọn cột 'Trạng thái' (Chưa chạy / Pass / Fail / Blocked / Skip); điền 'Kết quả thực tế', 'Người test', 'Ngày test'.",
        "Ưu tiên: P1 = bắt buộc pass trước khi phát hành; P2 = nên pass; P3 = có thể hoãn.",
        "Mã TC theo dạng TC-<PHÂN HỆ>-<TẦNG>-<SỐ>: API/INT = backend, E2E = giao diện, SEC/PERF = phi chức năng.",
    ], i + 1):
        ov.cell(row=j, column=1, value=txt).font = Font(name=FONT, size=10)
    ov.column_dimensions["A"].width = 34
    ov.column_dimensions["B"].width = 110
    wb.move_sheet("Tổng quan", offset=-wb.index(wb["Tổng quan"]))

    wb.calculation.fullCalcOnLoad = True  # Excel tự tính công thức khi mở
    wb.save(OUT)
    sys.stdout.reconfigure(encoding="utf-8")
    print(f"Đã ghi {OUT} — UC={len(ucs)}, bước={len(steps)}, TC v3={len(v3)}, TC v1={len(legacy)}")


if __name__ == "__main__":
    sys.exit(build())
